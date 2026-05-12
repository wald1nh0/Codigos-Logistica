import pandas as pd
import sys
import os
import tkinter as tk
from tkinter import filedialog
from openpyxl import load_workbook
from openpyxl.worksheet.table import Table, TableStyleInfo

# --- CONFIGURACIÓN ---
#ARCHIVO_CONSOLIDADO = "PUREBA.xlsx"
ARCHIVO_CONSOLIDADO = "Aqui va el nombre de tu archivo consolidado.xlsx"
HOJA_DIARIA = "Aqui va el nombre de la hoja diaria de donde vas a extraer la informacion (ej: EDITABLE)"

# VÍA 1: Tienen el ID largo (ej: 0999-12345-0001). Se limpia y va a Zipnova directo.
MARKET_ZIPNOVA = ['Dimarsa','Travel','Meli - ZipNova','Shopi - ZipNova','SAC - Bluex']
# VÍA 2: Tienen seguimiento Starken. Intentará consultar la web de Starken.
MARKET_STARKEN = ['Meli - Starken','Shopi - Starken','SAC - Starken']
# VÍA 3: Tienen seguimiento Blue. Van directo a la API de Blue Express.
MARKET_BLUE = ['Meli - Blue','Shopi - Blue']
MARKETS_VALIDOS = MARKET_ZIPNOVA + MARKET_STARKEN + MARKET_BLUE

print("--- INGESTOR DE PEDIDOS NUEVOS ---")

# 1. Cargar Consolidado Histórico (Si existe)
if os.path.exists(ARCHIVO_CONSOLIDADO):
    print(f"Cargando Base Histórica: {ARCHIVO_CONSOLIDADO}")
    df_base = pd.read_excel(ARCHIVO_CONSOLIDADO)
else:
    print("No se encontró base histórica. Se creará una nueva.")
    df_base = pd.DataFrame()

# 2. Seleccionar y Cargar Plantilla Diaria
print("Selecciona tu Plantilla Diaria (Excel)...")
root = tk.Tk(); root.withdraw()
path = filedialog.askopenfilename(title="Plantilla Diaria", filetypes=[("Excel", "*.xlsm *.xlsx")])
print(f"Extraccion Seleccionada: {path}")
if not path:
    print("Operación cancelada."); sys.exit()

try:
    df_nuevo = pd.read_excel(path, sheet_name=HOJA_DIARIA)
    # Filtramos solo las columnas que importan
    df_nuevo = df_nuevo[['Market','OPL','OC', 'SEG','SKU','PRODUCTO','Unidades','Bultos', 'Fecha Compra']]
    df_nuevo.rename(columns={'Seguimiento': 'SEG'}, inplace=True)
except Exception as e:
    print(f"Error leyendo la plantilla. Revisa las columnas: {e}")
    sys.exit()

# 3. Filtrar Markets Válidos
df_nuevo['Market'] = df_nuevo['Market'].astype(str).str.strip()
df_nuevo = df_nuevo[df_nuevo['Market'].isin(MARKETS_VALIDOS)]

# 4. Detectar Pedidos Nuevos (La Magia)
if not df_base.empty:
    existentes = df_base['SEG'].astype(str).tolist()
    # Filtramos los seguimientos que NO están en la base histórica
    df_agregar = df_nuevo[~df_nuevo['SEG'].astype(str).isin(existentes)].copy()
else:
    df_agregar = df_nuevo.copy()

nuevos_count = len(df_agregar)

if nuevos_count == 0:
    print("\n✅ No hay pedidos nuevos para agregar. La base ya está al día.")
    sys.exit()

print(f"\n🚀 Se encontraron {nuevos_count} pedidos nuevos. Agregando a la base...")

# 5. Preparar Columnas de Control para los Nuevos
columnas_control = ['Estado_Actual', 'Fecha_Recepcion_Courier', 'Fecha_Entrega_Real', 'Dias_Transcurridos', 'OTIF_Status', 'Comuna_Courier', 'Var_Medicion', 'Comentario']

for col in columnas_control:
    if col not in df_agregar.columns:
        df_agregar[col] = ""

# Ponemos un estado inicial para que el Actualizador sepa que debe buscarlos
df_agregar['Estado_Actual'] = "Por Consultar"
df_agregar['OTIF_Status'] = "Pendiente"

# 6. Unir y Guardar
df_final = pd.concat([df_base, df_agregar], ignore_index=True)

df_final.to_excel(ARCHIVO_CONSOLIDADO, index=False)

# 7. Aplicar Formato de Tabla Visual
try:
    wb = load_workbook(ARCHIVO_CONSOLIDADO)
    ws = wb.active
    
    tab = Table(displayName="TablaOTIF", ref=ws.dimensions)
    tab.tableStyleInfo = TableStyleInfo(name="TableStyleMedium9", showRowStripes=True)
    ws.add_table(tab)
    
    for col in ws.columns:
        ws.column_dimensions[col[0].column_letter].width = max(len(str(c.value)) for c in col if c.value) + 2

    wb.save(ARCHIVO_CONSOLIDADO)
    print("¡Datos guardados y formato aplicado!")
except Exception as e:
    print(f"Los datos se guardaron, pero hubo un error con el formato visual: {e}")

print(f"\n🎉 LISTO. Se agregaron {nuevos_count} filas. Ahora puedes correr el 'Actualizador_Estados.py'.")